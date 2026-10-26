"""Bounded regressions for the authored Tableau-to-Power-BI presentations."""

import contextlib
import importlib.util
import io
import posixpath
import re
import runpy
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import FrozenInstanceError
from itertools import pairwise
from pathlib import Path
from unittest import mock
from urllib.parse import urlsplit

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

try:
    import pptx as _python_pptx
except ModuleNotFoundError as exc:
    if exc.name != "pptx":
        raise
    PYTHON_PPTX_AVAILABLE = False
else:
    PYTHON_PPTX_AVAILABLE = True

if PYTHON_PPTX_AVAILABLE:
    from docs import generate_pptx as presentation_generator


WRAPPER_PATH = REPO_ROOT / "scripts" / "generate_presentation.py"
SAFE_MARGIN_EMU = int(0.30 * 914400)
RELATIONSHIP_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"

STALE_V28_RELEASE = re.compile(r"(?<![\w.])v28\.1\.1(?![\w.])", re.IGNORECASE)
STALE_V40_RELEASE = re.compile(r"(?<![\w.])v40\.0\.0(?![\w.])", re.IGNORECASE)

STALE_CLAIMS = (
    ("95% guarantee", re.compile(r"(?<!\d)95\s*%", re.IGNORECASE)),
    ("$0 guarantee", re.compile(r"\$\s*0(?:[.,]0+)?(?!\d)", re.IGNORECASE)),
    ("v28.1.1 release", STALE_V28_RELEASE),
    ("v40.0.0 release", STALE_V40_RELEASE),
    (
        "190-visual claim",
        re.compile(r"\b190\s+(?:visuals?|visuels?)\b", re.IGNORECASE),
    ),
)


def _named_shape_text(slide, name):
    matches = [shape for shape in slide.shapes if shape.name == name]
    if len(matches) != 1:
        raise AssertionError(
            f"expected one shape named {name!r}; found {len(matches)}"
        )
    return matches[0].text


def _slide_text(slide):
    return "\n".join(
        shape.text for shape in slide.shapes if shape.has_text_frame
    )


def _notes_text(slide):
    return slide.notes_slide.notes_text_frame.text


def _find_slide(presentation, title):
    for slide in presentation.slides:
        if _named_shape_text(slide, "Slide title").strip() == title:
            return slide
    raise AssertionError(f"slide title not found: {title!r}")


def _all_deck_text(presentation):
    parts = []
    for slide in presentation.slides:
        parts.extend((_slide_text(slide), _notes_text(slide)))
    return "\n".join(parts)


def _normalized(text):
    return re.sub(r"\s+", " ", text).strip().casefold()


@unittest.skipUnless(
    PYTHON_PPTX_AVAILABLE,
    "optional python-pptx dependency is not installed",
)
class TestDocumentationPresentations(unittest.TestCase):
    """Generate the real decks once, then inspect their public contracts."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._temporary_directory = tempfile.TemporaryDirectory(
            prefix="documentation-presentations-"
        )
        output_dir = Path(cls._temporary_directory.name)
        cls.overview_path = output_dir / "overview.pptx"
        cls.guide_path = output_dir / "guide.pptx"
        with contextlib.redirect_stdout(io.StringIO()):
            cls.overview_result = presentation_generator.generate_presentation(
                cls.overview_path
            )
            cls.guide_result = presentation_generator.generate_technical_guide(
                cls.guide_path
            )
        cls.overview = _python_pptx.Presentation(str(cls.overview_path))
        cls.guide = _python_pptx.Presentation(str(cls.guide_path))

    @classmethod
    def tearDownClass(cls):
        cls._temporary_directory.cleanup()
        super().tearDownClass()

    def test_business_case_defaults_are_explicit_and_immutable(self):
        case = presentation_generator.BusinessCase(
            assisted_hours_per_report=5
        )
        self.assertEqual(
            (
                case.report_count,
                case.manual_hours_per_report,
                case.setup_hours,
                case.hourly_rate,
                case.hours_per_day,
            ),
            (100, 16, 80, 100, 8),
        )
        with self.assertRaises(FrozenInstanceError):
            case.setup_hours = 81

    def test_scenarios_preserve_the_published_calculations(self):
        expected = {
            "Prudent": (880, 720, 45, 90, 72000),
            "Central": (580, 1020, 63.75, 127.5, 102000),
            "Ambitieux": (380, 1220, 76.25, 152.5, 122000),
        }
        self.assertEqual(
            [name for name, _ in presentation_generator.SCENARIOS],
            list(expected),
        )
        for name, case in presentation_generator.SCENARIOS:
            with self.subTest(scenario=name):
                assisted, saved, percent, person_days, value = expected[name]
                self.assertEqual(case.assisted_hours, assisted)
                self.assertEqual(case.saved_hours, saved)
                self.assertAlmostEqual(case.saved_percent, percent)
                self.assertAlmostEqual(case.saved_person_days, person_days)
                self.assertEqual(case.labor_value, value)

    def test_business_case_keeps_negative_savings_as_a_loss(self):
        case = presentation_generator.BusinessCase(
            assisted_hours_per_report=20
        )
        self.assertEqual(case.saved_hours, -480)
        self.assertEqual(case.saved_percent, -30)
        self.assertEqual(case.saved_person_days, -60)
        self.assertEqual(case.labor_value, -48000)

    def test_report_count_rejects_booleans_nonintegers_and_nonpositive_values(self):
        for report_count in (True, False, 0, -1, 1.5, "100", None):
            with self.subTest(report_count=report_count), self.assertRaises(
                ValueError
            ):
                presentation_generator.BusinessCase(
                    assisted_hours_per_report=5,
                    report_count=report_count,
                )

    def test_hours_and_rate_reject_nonfinite_negative_and_boolean_values(self):
        fields = (
            "assisted_hours_per_report",
            "manual_hours_per_report",
            "setup_hours",
            "hourly_rate",
            "hours_per_day",
        )
        invalid_values = (-1, float("nan"), float("inf"), -float("inf"), True)
        for field in fields:
            for value in invalid_values:
                with self.subTest(field=field, value=value), self.assertRaises(
                    ValueError
                ):
                    arguments = {"assisted_hours_per_report": 5}
                    arguments[field] = value
                    presentation_generator.BusinessCase(**arguments)

    def test_manual_hours_and_day_length_must_be_positive(self):
        for field in ("manual_hours_per_report", "hours_per_day"):
            with self.subTest(field=field), self.assertRaises(ValueError):
                presentation_generator.BusinessCase(
                    assisted_hours_per_report=5,
                    **{field: 0},
                )

    def test_zero_setup_effort_and_zero_rate_are_valid_inputs(self):
        case = presentation_generator.BusinessCase(
            assisted_hours_per_report=0,
            setup_hours=0,
            hourly_rate=0,
        )
        self.assertEqual(case.assisted_hours, 0)
        self.assertEqual(case.saved_hours, 1600)
        self.assertEqual(case.labor_value, 0)

    def test_standard_deck_filenames_are_stable(self):
        self.assertEqual(
            presentation_generator.OVERVIEW_NAME,
            "Tableau_to_PowerBI_Migration.pptx",
        )
        self.assertEqual(
            presentation_generator.GUIDE_NAME,
            "Tableau_to_PowerBI_Migration_Guide.pptx",
        )

    def test_public_generators_return_the_original_path_object(self):
        self.assertIs(self.overview_result, self.overview_path)
        self.assertIs(self.guide_result, self.guide_path)
        self.assertTrue(self.overview_path.is_file())
        self.assertTrue(self.guide_path.is_file())

    def test_package_metadata_is_project_authored_and_dated_to_the_edition(self):
        for deck_name, presentation in (
            ("overview", self.overview),
            ("guide", self.guide),
        ):
            properties = presentation.core_properties
            with self.subTest(deck=deck_name):
                self.assertEqual(properties.author, "TableauToPowerBI project")
                self.assertEqual(properties.last_modified_by, properties.author)
                self.assertEqual(
                    properties.created, presentation_generator.EDITION_TIMESTAMP
                )
                self.assertEqual(
                    properties.modified, presentation_generator.EDITION_TIMESTAMP
                )
                self.assertEqual(
                    properties.comments, "Generated by docs/generate_pptx.py"
                )

    def test_internal_generate_returns_the_exact_passed_path(self):
        with tempfile.TemporaryDirectory(
            prefix="documentation-generate-contract-"
        ) as temporary_directory:
            destination = Path(temporary_directory) / "contract.pptx"
            with (
                contextlib.redirect_stdout(io.StringIO()),
                mock.patch.object(presentation_generator, "_overview") as build,
            ):
                returned = presentation_generator._generate(
                    destination,
                    technical=False,
                )
            self.assertIs(returned, destination)
            self.assertTrue(destination.is_file())
            build.assert_called_once()

    def test_overview_slide_count_and_named_bookends(self):
        self.assertEqual(len(self.overview.slides), 15)
        self.assertEqual(
            _named_shape_text(self.overview.slides[0], "Deck title"),
            "Accélérer la migration.\nGarder le contrôle.",
        )
        self.assertEqual(
            _named_shape_text(self.overview.slides[-1], "Closing title"),
            "Commencer petit.\nDécider sur des preuves.",
        )
        for slide in list(self.overview.slides)[1:-1]:
            self.assertTrue(_named_shape_text(slide, "Slide title").strip())

    def test_guide_slide_count_and_named_bookends(self):
        self.assertEqual(len(self.guide.slides), 23)
        self.assertEqual(
            _named_shape_text(self.guide.slides[0], "Deck title"),
            "Une architecture ouverte,\ndes contrôles explicites.",
        )
        self.assertEqual(
            _named_shape_text(self.guide.slides[-1], "Closing title"),
            "Un socle généré.\nUne acceptation à démontrer.",
        )
        for slide in list(self.guide.slides)[1:-1]:
            self.assertTrue(_named_shape_text(slide, "Slide title").strip())

    def test_overview_key_sections_follow_a_coherent_order(self):
        titles = [
            _named_shape_text(slide, "Slide title").strip()
            for slide in self.overview.slides
        ]
        ordered_titles = (
            "Hyper ou source native : une règle simple",
            "Une preuve statique, pas une promesse de parité",
            "Comparer le même périmètre de travail",
            "Trois scénarios, pas une promesse",
            "Valoriser la capacité potentiellement libérée",
            "Du gain brut au coût total du projet",
            "Sources, hypothèses et périmètre",
        )
        for earlier, later in pairwise(ordered_titles):
            with self.subTest(order=(earlier, later)):
                self.assertLess(titles.index(earlier), titles.index(later))

    def test_guide_key_sections_follow_a_coherent_order(self):
        titles = [
            _named_shape_text(slide, "Slide title").strip()
            for slide in self.guide.slides
        ]
        ordered_titles = (
            "Trois routes de sortie, trois contrats",
            "PBIP : un projet lisible, pas une boîte noire",
            "PBIR : conserver l'intention visuelle",
            "De la formule Tableau au bon contexte",
            "Prep autonome : M, sources et lignage",
            "Fabric : un scaffold à configurer et déployer",
            "Une preuve statique, pas une promesse de parité",
        )
        for earlier, later in pairwise(ordered_titles):
            with self.subTest(order=(earlier, later)):
                self.assertLess(titles.index(earlier), titles.index(later))

    def test_slide_notes_carry_edition_provenance_and_source_sections(self):
        expected_header = f"Édition du {presentation_generator.EDITION_DATE}."
        expected_footer = (
            "Sources de contenu (checkout de travail ; pas une certification) :"
        )
        for deck_name, presentation in (
            ("overview", self.overview),
            ("guide", self.guide),
        ):
            citations = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                notes = _notes_text(slide)
                with self.subTest(deck=deck_name, slide=slide_number):
                    self.assertIn(expected_header, notes)
                    self.assertIn(expected_footer, notes)
                if "/blob/main/" in notes:
                    citations.append(notes)
            self.assertTrue(citations, f"{deck_name} has no speaker-note citations")
            self.assertIn(
                "docs/ARCHITECTURE.md",
                "\n".join(citations),
            )

    def test_overview_business_case_notes_disclose_illustrative_assumptions(self):
        comparison = _notes_text(
            _find_slide(
                self.overview,
                "Comparer le même périmètre de travail",
            )
        )
        for disclosure in (
            "80 heures de setup sont ajoutées une seule fois",
            "Jours-personnes de 8 heures, pas un délai calendaire",
            "n'établit pas une économie de trésorerie",
        ):
            with self.subTest(disclosure=disclosure):
                self.assertIn(disclosure, comparison)

        scenarios = _notes_text(
            _find_slide(self.overview, "Trois scénarios, pas une promesse")
        )
        self.assertIn("Prudent : 80 + 100*8 = 880 h", scenarios)
        self.assertIn("Central : 80 + 100*5 = 580 h", scenarios)
        self.assertIn("Ambitieux : 80 + 100*3 = 380 h", scenarios)
        self.assertIn(
            "Aucune de ces valeurs n'est un résultat de mesure terrain",
            scenarios,
        )

    def test_overview_speaker_notes_include_each_scenario_result(self):
        notes = _notes_text(
            _find_slide(
                self.overview,
                "Valoriser la capacité potentiellement libérée",
            )
        )
        expected_metrics = {
            "Prudent": (720, 45, 90, 72000),
            "Central": (1020, 63.75, 127.5, 102000),
            "Ambitieux": (1220, 76.25, 152.5, 122000),
        }
        for name, (saved, percent, person_days, value) in expected_metrics.items():
            line = next(
                line for line in notes.splitlines() if line.startswith(f"{name}:")
            )
            self.assertIn(f"saved={saved}h", line)
            percent_match = re.search(r"\(([-+]?\d+(?:\.\d+)?)%\)", line)
            self.assertIsNotNone(percent_match, line)
            self.assertAlmostEqual(float(percent_match.group(1)), percent)
            self.assertIn(f"person_days={person_days}", line)
            self.assertIn(f"labor_value={value} EUR", line)

    def test_overview_commercial_copy_clarifies_avoided_effort_and_value(self):
        savings_slide = _find_slide(
            self.overview,
            "Valoriser la capacité potentiellement libérée",
        )
        percentage_labels = {
            shape.text
            for shape in savings_slide.shapes
            if shape.has_text_frame and "%" in shape.text
        }
        self.assertEqual(
            percentage_labels,
            {
                "45 % de charge évitée",
                "63,75 % de charge évitée",
                "76,25 % de charge évitée",
            },
        )

        economic_slide = _slide_text(
            _find_slide(self.overview, "Du gain brut au coût total du projet")
        )
        for clarification in (
            "Cas central illustratif",
            "Capacité valorisée − coûts",
            "Ni économie de trésorerie garantie, ni ROI mesuré.",
        ):
            with self.subTest(clarification=clarification):
                self.assertIn(clarification, economic_slide)

    def _assert_source_policy(self, presentation):
        slide = _find_slide(
            presentation,
            "Hyper ou source native : une règle simple",
        )
        visible = _normalized(_slide_text(slide))
        notes = _normalized(_notes_text(slide))
        for phrase in (
            "≤ 10 mib",
            "csv + import",
            "> 10 mib",
            "parquet + import",
        ):
            with self.subTest(surface="visible", phrase=phrase):
                self.assertIn(phrase, visible)
        for phrase in (
            "10 * 1024 * 1024 octets",
            "strictement au-dessus",
            "parquet natif + import",
            "sdk hyper optionnel requis",
            "directquery si le connecteur le permet",
            "sinon import",
            "l'utilisateur peut forcer import",
            "aucune extraction/copie de lignes par l'outil",
            "excel réellement embarqué dans le twbx",
            "copié sans modification dans data/",
            "lié par chemin absolu",
            "ce n'est pas une extraction de lignes",
            "ni les jobs téléversés ni les modèles partagés",
        ):
            with self.subTest(surface="speaker notes", phrase=phrase):
                self.assertIn(phrase, notes)
        self.assertIn("aucune extraction des données natives", visible)
        self.assertIn("excel embarqué : copié tel quel dans data/", visible)

    def test_overview_source_policy_covers_hyper_and_native_sources(self):
        self._assert_source_policy(self.overview)

    def test_guide_source_policy_covers_hyper_and_native_sources(self):
        self._assert_source_policy(self.guide)

    def test_guide_pbip_anatomy_states_the_packaged_excel_exception(self):
        slide = _find_slide(
            self.guide,
            "PBIP : un projet lisible, pas une boîte noire",
        )
        visible = _normalized(_slide_text(slide))
        for phrase in (
            "sources natives externes",
            "sauf excel embarqué copié dans data/",
            "l'import charge à l'actualisation",
        ):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, visible)
        self.assertNotIn("restent externes", visible)

    def test_static_evidence_is_dated_and_separated_from_runtime_claims(self):
        for deck_name, presentation in (
            ("overview", self.overview),
            ("guide", self.guide),
        ):
            slide = _find_slide(
                presentation,
                "Une preuve statique, pas une promesse de parité",
            )
            visible_and_notes = _normalized(
                _slide_text(slide) + "\n" + _notes_text(slide)
            )
            self.assertIn(presentation_generator.EVIDENCE_DATE, visible_and_notes)
            for caveat in (
                "statique",
                "historique",
                "ci courante",
                "moteur dax",
                "desktop",
                "refresh",
                "gains de temps",
            ):
                with self.subTest(deck=deck_name, caveat=caveat):
                    self.assertIn(caveat, visible_and_notes)

    def test_stale_release_patterns_distinguish_format_and_release_versions(self):
        self.assertIsNone(STALE_V28_RELEASE.search("PBIR v4.0"))
        self.assertIsNone(STALE_V40_RELEASE.search("PBIR v4.0"))
        self.assertEqual(STALE_V28_RELEASE.search("v28.1.1").group(0), "v28.1.1")
        self.assertEqual(STALE_V40_RELEASE.search("v40.0.0").group(0), "v40.0.0")

    def test_rendered_decks_exclude_stale_numeric_and_version_claims(self):
        violations = []
        for deck_name, presentation in (
            ("overview", self.overview),
            ("guide", self.guide),
        ):
            deck_text = _all_deck_text(presentation)
            for label, pattern in STALE_CLAIMS:
                matches = [match.group(0) for match in pattern.finditer(deck_text)]
                violations.extend(
                    f"{deck_name}: {label} ({match!r})" for match in matches
                )
        self.assertEqual([], violations, "\n".join(violations))

    def _assert_slide_geometry(self, presentation):
        slide_width = presentation.slide_width
        slide_height = presentation.slide_height
        line_type = _python_pptx.enum.shapes.MSO_SHAPE_TYPE.LINE
        for slide_number, slide in enumerate(presentation.slides, start=1):
            self.assertTrue(slide.shapes, f"slide {slide_number} is empty")
            for shape in slide.shapes:
                label = f"slide {slide_number}, shape {shape.name!r}"
                self.assertGreaterEqual(shape.width, 0, label)
                self.assertGreaterEqual(shape.height, 0, label)
                if shape.shape_type == line_type:
                    self.assertGreater(
                        shape.width + shape.height,
                        0,
                        label,
                    )
                else:
                    self.assertGreater(shape.width, 0, label)
                    self.assertGreater(shape.height, 0, label)
                self.assertGreaterEqual(shape.left, SAFE_MARGIN_EMU, label)
                self.assertGreaterEqual(shape.top, SAFE_MARGIN_EMU, label)
                self.assertLessEqual(
                    shape.left + shape.width,
                    slide_width - SAFE_MARGIN_EMU,
                    label,
                )
                self.assertLessEqual(
                    shape.top + shape.height,
                    slide_height - SAFE_MARGIN_EMU,
                    label,
                )

    def test_overview_shapes_have_positive_geometry_and_safe_margins(self):
        self._assert_slide_geometry(self.overview)

    def test_guide_shapes_have_positive_geometry_and_safe_margins(self):
        self._assert_slide_geometry(self.guide)

    def _assert_native_shapes_and_public_logo(self, presentation, deck_path):
        picture_type = _python_pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE
        auto_shape_type = _python_pptx.enum.shapes.MSO_SHAPE_TYPE.AUTO_SHAPE
        pictures = [
            (slide, shape)
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.shape_type == picture_type
        ]
        self.assertEqual(len(pictures), 1)
        self.assertIs(pictures[0][0], presentation.slides[0])
        for slide_number, slide in enumerate(presentation.slides, start=1):
            self.assertTrue(
                any(shape.has_text_frame for shape in slide.shapes),
                f"slide {slide_number} has no editable text",
            )
            self.assertTrue(
                any(shape.shape_type == auto_shape_type for shape in slide.shapes),
                f"slide {slide_number} has no native editable shape",
            )

        logo_path = (
            presentation_generator.DOCS_DIR
            / "images"
            / "logo-migration-dashboards.png"
        )
        with zipfile.ZipFile(deck_path) as archive:
            media_parts = sorted(
                name
                for name in archive.namelist()
                if name.startswith("ppt/media/")
            )
            self.assertEqual(len(media_parts), 1)
            self.assertEqual(archive.read(media_parts[0]), logo_path.read_bytes())

    def test_overview_uses_editable_shapes_and_only_the_public_logo_image(self):
        self._assert_native_shapes_and_public_logo(self.overview, self.overview_path)

    def test_guide_uses_editable_shapes_and_only_the_public_logo_image(self):
        self._assert_native_shapes_and_public_logo(self.guide, self.guide_path)

    def _assert_package_relationships(self, deck_path):
        external_hyperlinks = []
        with zipfile.ZipFile(deck_path) as archive:
            members = set(archive.namelist())
            self.assertIsNone(archive.testzip())
            self.assertIn("[Content_Types].xml", members)
            self.assertIn("ppt/presentation.xml", members)
            relationship_parts = sorted(
                name for name in members if name.endswith(".rels")
            )
            self.assertTrue(relationship_parts)

            for relationship_part in relationship_parts:
                root = ET.fromstring(archive.read(relationship_part))
                for relationship in root.findall(
                    f"{RELATIONSHIP_NS}Relationship"
                ):
                    target = relationship.attrib["Target"]
                    relationship_type = relationship.attrib["Type"]
                    target_mode = relationship.attrib.get(
                        "TargetMode",
                        "",
                    ).casefold()
                    if target_mode == "external":
                        self.assertTrue(
                            relationship_type.endswith("/hyperlink"),
                            f"non-hyperlink external relationship: {relationship_part}",
                        )
                        parsed_target = urlsplit(target)
                        self.assertEqual(parsed_target.scheme, "https")
                        self.assertEqual(parsed_target.netloc.casefold(), "github.com")
                        self.assertTrue(
                            target.startswith(
                                f"{presentation_generator.REPOSITORY}/blob/main/"
                            ),
                            target,
                        )
                        external_hyperlinks.append(target)
                        continue

                    if relationship_part == "_rels/.rels":
                        relationship_directory = ""
                    else:
                        relationship_directory = relationship_part.split(
                            "/_rels/",
                            maxsplit=1,
                        )[0]
                    if target.startswith("/"):
                        resolved_target = posixpath.normpath(target.lstrip("/"))
                    else:
                        resolved_target = posixpath.normpath(
                            posixpath.join(relationship_directory, target)
                        )
                    self.assertFalse(
                        resolved_target == ".."
                        or resolved_target.startswith("../"),
                        f"relationship escapes package: {relationship_part} -> {target}",
                    )
                    self.assertIn(
                        resolved_target,
                        members,
                        f"missing target: {relationship_part} -> {target}",
                    )
        return external_hyperlinks

    def test_overview_zip_crc_and_internal_relationship_targets(self):
        self._assert_package_relationships(self.overview_path)

    def test_guide_zip_crc_and_internal_relationship_targets(self):
        self._assert_package_relationships(self.guide_path)

    def test_external_ooxml_relationships_are_public_hyperlinks_only(self):
        overview_links = self._assert_package_relationships(self.overview_path)
        guide_links = self._assert_package_relationships(self.guide_path)
        self.assertTrue(overview_links)
        self.assertTrue(guide_links)

    def test_documentation_cli_defaults_to_the_overview(self):
        with tempfile.TemporaryDirectory(prefix="presentation-cli-") as folder:
            with (
                mock.patch.object(
                    sys,
                    "argv",
                    ["generate_pptx.py", "--output-dir", folder],
                ),
                mock.patch.object(
                    presentation_generator,
                    "generate_presentation",
                ) as generate_overview,
                mock.patch.object(
                    presentation_generator,
                    "generate_technical_guide",
                ) as generate_guide,
            ):
                presentation_generator.main()
            generate_overview.assert_called_once_with(
                Path(folder) / presentation_generator.OVERVIEW_NAME
            )
            generate_guide.assert_not_called()

    def test_cli_output_and_short_alias_dispatch_a_single_deck(self):
        for option in ("--output", "-o"):
            with (
                self.subTest(option=option),
                tempfile.TemporaryDirectory(
                    prefix="presentation-cli-file-"
                ) as folder,
            ):
                destination = Path(folder) / "custom.pptx"
                with (
                    mock.patch.object(
                        sys,
                        "argv",
                        [
                            "generate_pptx.py",
                            "--deck",
                            "guide",
                            option,
                            str(destination),
                        ],
                    ),
                    mock.patch.object(
                        presentation_generator,
                        "generate_presentation",
                    ) as generate_overview,
                    mock.patch.object(
                        presentation_generator,
                        "generate_technical_guide",
                    ) as generate_guide,
                ):
                    presentation_generator.main()
                generate_guide.assert_called_once_with(destination)
                generate_overview.assert_not_called()

    def test_cli_both_uses_standard_filenames_in_the_output_directory(self):
        with tempfile.TemporaryDirectory(prefix="presentation-cli-both-") as folder:
            with (
                mock.patch.object(
                    sys,
                    "argv",
                    [
                        "generate_pptx.py",
                        "--deck",
                        "both",
                        "--output-dir",
                        folder,
                    ],
                ),
                mock.patch.object(
                    presentation_generator,
                    "generate_presentation",
                ) as generate_overview,
                mock.patch.object(
                    presentation_generator,
                    "generate_technical_guide",
                ) as generate_guide,
            ):
                presentation_generator.main()
            generate_overview.assert_called_once_with(
                Path(folder) / presentation_generator.OVERVIEW_NAME
            )
            generate_guide.assert_called_once_with(
                Path(folder) / presentation_generator.GUIDE_NAME
            )

    def test_cli_rejects_file_and_directory_outputs_together(self):
        error_output = io.StringIO()
        with (
            tempfile.TemporaryDirectory(prefix="presentation-cli-xor-") as folder,
            mock.patch.object(
                sys,
                "argv",
                [
                    "generate_pptx.py",
                    "--output",
                    str(Path(folder) / "custom.pptx"),
                    "--output-dir",
                    folder,
                ],
            ),
            mock.patch.object(
                presentation_generator,
                "generate_presentation",
            ) as generate_overview,
            mock.patch.object(
                presentation_generator,
                "generate_technical_guide",
            ) as generate_guide,
            contextlib.redirect_stderr(error_output),
            self.assertRaises(SystemExit) as error,
        ):
            presentation_generator.main()
        self.assertEqual(error.exception.code, 2)
        self.assertIn("not allowed with argument", error_output.getvalue())
        generate_overview.assert_not_called()
        generate_guide.assert_not_called()

    def test_cli_rejects_a_single_output_path_for_both_decks(self):
        error_output = io.StringIO()
        with (
            tempfile.TemporaryDirectory(
                prefix="presentation-cli-both-file-"
            ) as folder,
            mock.patch.object(
                sys,
                "argv",
                [
                    "generate_pptx.py",
                    "--deck",
                    "both",
                    "--output",
                    str(Path(folder) / "custom.pptx"),
                ],
            ),
            mock.patch.object(
                presentation_generator,
                "generate_presentation",
            ) as generate_overview,
            mock.patch.object(
                presentation_generator,
                "generate_technical_guide",
            ) as generate_guide,
            contextlib.redirect_stderr(error_output),
            self.assertRaises(SystemExit) as error,
        ):
            presentation_generator.main()
        self.assertEqual(error.exception.code, 2)
        self.assertIn(
            "--output is for a single deck",
            error_output.getvalue(),
        )
        generate_overview.assert_not_called()
        generate_guide.assert_not_called()

    def test_wrapper_keeps_the_no_argument_guide_builder(self):
        spec = importlib.util.spec_from_file_location(
            "_documentation_presentation_wrapper",
            WRAPPER_PATH,
        )
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        wrapper = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(wrapper)

        with mock.patch.object(
            wrapper,
            "generate_technical_guide",
            return_value=object(),
        ) as generate_guide:
            result = wrapper.build_presentation()
        self.assertIs(result, generate_guide.return_value)
        generate_guide.assert_called_once_with(
            wrapper.DOCS_DIR / wrapper.GUIDE_NAME
        )

    def test_wrapper_cli_selects_the_guide_as_its_default(self):
        with mock.patch.object(presentation_generator, "main") as main:
            runpy.run_path(str(WRAPPER_PATH), run_name="__main__")
        main.assert_called_once_with(default_deck="guide")


if __name__ == "__main__":
    unittest.main()

#!/usr/bin/env python3
"""Build the French overview/business-case and technical documentation decks.

Optional authoring dependency: python-pptx. Core migration does not import this
module. Both decks use native, editable PowerPoint shapes and a public logo.

    python docs\\generate_pptx.py --deck both
    python docs\\generate_pptx.py --deck guide --output guide.pptx
"""

import argparse
import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

DOCS_DIR = Path(__file__).resolve().parent
OVERVIEW_NAME = "Tableau_to_PowerBI_Migration.pptx"
GUIDE_NAME = "Tableau_to_PowerBI_Migration_Guide.pptx"
EDITION_DATE = "26.10.2026"
EDITION_TIMESTAMP = datetime.strptime(EDITION_DATE, "%d.%m.%Y")
EVIDENCE_DATE = "23.10.2026"
REPOSITORY = "https://github.com/cyphou/TableauToPowerBI"
SLIDE_WIDTH = Inches(13.333333)
SLIDE_HEIGHT = Inches(7.5)

INK = "182630"
MUTED = "52616C"
PAPER = "F5F7F8"
WHITE = "FFFFFF"
TEAL = "007F79"
TEAL_LIGHT = "E0F1EF"
YELLOW = "F2C811"
ORANGE = "C85A17"
ORANGE_LIGHT = "FCEDDF"
BLUE = "1266A0"
GRAY = "DDE4E8"
NIGHT = "0D1D26"
NIGHT_CARD = "20353F"
LIGHT_TEXT = "C9D8DF"
FONT = "Aptos"
FONT_TITLE = "Aptos Display"

REFERENCES = {
    "architecture": ("Architecture et modes de source", "docs/ARCHITECTURE.md"),
    "features": ("Fonctionnalités et preuves datées", "README.md"),
    "mapping": ("Inventaire des correspondances", "docs/MAPPING_REFERENCE.md"),
    "prep": ("Périmètre Tableau Prep", "docs/TABLEAU_PREP_TO_POWERQUERY_REFERENCE.md"),
    "agents": ("Propriété et handoffs des agents", "docs/AGENTS.md"),
    "tools": ("Surface MCP et assistance IA", "docs/AGENT_SURFACE.md"),
    "quality": ("Contrôles et critères de recette", "docs/MIGRATION_CHECKLIST.md"),
    "limits": ("Limites et approximations", "docs/KNOWN_LIMITATIONS.md"),
    "deploy": ("Publication séparément autorisée", "docs/DEPLOYMENT_GUIDE.md"),
    "cli": ("Commandes et exploitation", "docs/USER_MANUAL.md"),
    "authoring": ("Validation Microsoft optionnelle", "docs/POWERBI_PROJECT_GUIDE.md"),
    "tests": ("Tests et portée des résultats", "tests/README.md"),
}


@dataclass(frozen=True)
class BusinessCase:
    """Illustrative workload model, not a measured benchmark or cash forecast."""

    assisted_hours_per_report: float
    report_count: int = 100
    manual_hours_per_report: float = 16
    setup_hours: float = 80
    hourly_rate: float = 100
    hours_per_day: float = 8

    def __post_init__(self):
        if (
            isinstance(self.report_count, bool)
            or not isinstance(self.report_count, int)
            or self.report_count <= 0
        ):
            raise ValueError("report_count must be a positive integer")
        for name in (
            "assisted_hours_per_report", "manual_hours_per_report",
            "setup_hours", "hourly_rate", "hours_per_day",
        ):
            value = getattr(self, name)
            if isinstance(value, bool) or not math.isfinite(value) or value < 0:
                raise ValueError(f"{name} must be finite and nonnegative")
        if self.manual_hours_per_report == 0 or self.hours_per_day == 0:
            raise ValueError("manual_hours_per_report and hours_per_day must be positive")

    @property
    def manual_hours(self) -> float:
        return self.report_count * self.manual_hours_per_report

    @property
    def assisted_hours(self) -> float:
        return self.setup_hours + self.report_count * self.assisted_hours_per_report

    @property
    def saved_hours(self) -> float:
        return self.manual_hours - self.assisted_hours

    @property
    def saved_percent(self) -> float:
        return self.saved_hours / self.manual_hours * 100

    @property
    def saved_person_days(self) -> float:
        return self.saved_hours / self.hours_per_day

    @property
    def labor_value(self) -> float:
        return self.saved_hours * self.hourly_rate


SCENARIOS = (
    ("Prudent", BusinessCase(8)),
    ("Central", BusinessCase(5)),
    ("Ambitieux", BusinessCase(3)),
)


def _fr(value: float, decimals: int = 0) -> str:
    return f"{value:,.{decimals}f}".replace(",", " ").replace(".", ",")


def _color(value):
    return RGBColor.from_string(value)


def _rect(slide, x, y, w, h, fill, *, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def _text(slide, x, y, w, h, text, size: float = 16, color=INK, *,
          bold=False, align=PP_ALIGN.LEFT, middle=False, name=None, font=FONT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if middle else MSO_ANCHOR.TOP
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(5)
        paragraph.line_spacing = 1.08
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _color(color)
    return shape


def _arrow(slide, x, y, w=0.22, h=0.22, color=MUTED):
    return _rect_shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, color)


def _rect_shape(slide, shape_type, x, y, w, h, fill):
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    shape.line.fill.background()
    return shape


def _line(slide, x1, y1, x2, y2, color=GRAY, width=2):
    line = slide.shapes.add_connector(
        1, Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    line.line.color.rgb = _color(color)
    line.line.width = Pt(width)
    return line


def _tag(slide, x, y, w, label, fill=TEAL_LIGHT, color=TEAL):
    _rect(slide, x, y, w, 0.35, fill, rounded=True)
    _text(slide, x + 0.10, y + 0.04, w - 0.20, 0.24,
          label, 11, color, bold=True)


def _panel(slide, x, y, w, h, title, body, *, fill=WHITE, color=INK,
           accent=TEAL, number=None):
    _rect(slide, x, y, w, h, fill, rounded=True)
    if not body:
        _text(slide, x + 0.23, y + 0.20, w - 0.46, h - 0.40,
              title, 21, accent, bold=True, middle=True)
        return
    compact = h <= 1.65
    heading_x = x + 0.23
    heading_w = w - 0.46
    if number is not None:
        _rect_shape(slide, MSO_SHAPE.OVAL, x + 0.23, y + 0.22, 0.38, 0.38, accent)
        _text(slide, x + 0.23, y + 0.255, 0.38, 0.24,
              str(number), 12, WHITE, bold=True, align=PP_ALIGN.CENTER)
        heading_x += 0.53
        heading_w -= 0.53
    _text(slide, heading_x, y + 0.18, heading_w, 0.70 if not compact else 0.43,
          title, 19, accent, bold=True)
    body_top = 0.66 if compact else 0.97
    _text(slide, x + 0.23, y + body_top, w - 0.46, h - body_top - 0.10,
          body, 14.5 if compact else 16, color)


def _node(slide, x, y, w, title, body, *, fill=INK, color=WHITE, h=1.38):
    _rect(slide, x, y, w, h, fill, rounded=True)
    _text(slide, x + 0.18, y + 0.19, w - 0.36, 0.44,
          title, 19, color, bold=True, align=PP_ALIGN.CENTER)
    _text(slide, x + 0.18, y + 0.74, w - 0.36, h - 0.81,
          body, 14, color, align=PP_ALIGN.CENTER)


def _banner(slide, text, *, y=6.04, fill=INK, color=WHITE):
    _rect(slide, 0.65, y, 12.03, 0.47, fill, rounded=True)
    _text(slide, 0.86, y + 0.10, 11.61, 0.28, text, 13, color, bold=True)


def _notes(slide, text, sources):
    references = "\n".join(
        f"{REFERENCES[key][0]}: {REPOSITORY}/blob/main/{REFERENCES[key][1]}"
        for key in sources
    )
    slide.notes_slide.notes_text_frame.text = (
        f"Édition du {EDITION_DATE}. Les preuves historiques sont datées séparément.\n"
        f"{text}\n\nSources de contenu (checkout de travail ; pas une certification) :\n"
        f"{references}"
    )


def _slide(prs, title, subtitle, section, *, sources=(), note="", dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _color(NIGHT if dark else PAPER)
    _text(slide, 0.65, 0.51, 11.8, 0.22, section.upper(), 10.5,
          YELLOW if dark else TEAL, bold=True)
    _text(slide, 0.65, 0.94, 12.03, 0.75, title, 36,
          WHITE if dark else INK, bold=True, font=FONT_TITLE, name="Slide title")
    if subtitle:
        _text(slide, 0.65, 1.83, 12.03, 0.54, subtitle, 15,
              LIGHT_TEXT if dark else MUTED)
    number = len(prs.slides)
    _text(slide, 0.65, 6.78, 11.1, 0.20,
          f"Tableau to Power BI  |  Édition {EDITION_DATE}  |  {section}",
          10, LIGHT_TEXT if dark else MUTED)
    _text(slide, 12.03, 6.76, 0.65, 0.23, f"{number:02d}", 11,
          LIGHT_TEXT if dark else MUTED, align=PP_ALIGN.RIGHT)
    _notes(slide, note, sources)
    return slide


def _cover(prs, *, technical=False):
    title = (
        "Une architecture ouverte,\ndes contrôles explicites."
        if technical else "Accélérer la migration.\nGarder le contrôle."
    )
    slide = _slide(prs, "", "", "Guide technique" if technical else "Vision & valeur",
                   sources=("architecture", "quality"), dark=True)
    _text(slide, 0.65, 1.36, 8.2, 0.45, "TABLEAU  →  POWER BI", 22, YELLOW, bold=True)
    _text(slide, 0.65, 2.22, 8.5, 1.70, title, 42, WHITE,
          bold=True, font=FONT_TITLE, name="Deck title")
    _text(
        slide, 0.65, 4.35, 8.3, 0.82,
        "Architecture, workflows, fonctionnalités et exploitation."
        if technical else
        "Industrialiser les tâches répétitives ;\nréserver l'expertise à la validation métier.",
        19, LIGHT_TEXT,
    )
    logo = DOCS_DIR / "images" / "logo-migration-dashboards.png"
    slide.shapes.add_picture(str(logo), Inches(9.55), Inches(2.26), width=Inches(2.65))
    _tag(slide, 0.65, 5.79, 3.5, "PROJETS ÉDITABLES", YELLOW, INK)
    _tag(slide, 4.45, 5.79, 3.7, "DÉCISION HUMAINE", NIGHT_CARD, WHITE)
    _notes(slide, "Support en français. Aucun résultat client ni gain de temps mesuré "
           "n'est présenté. Les schémas sont des formes PowerPoint éditables.",
           ("architecture", "quality"))


def _pipeline(prs, section):
    slide = _slide(
        prs, "Deux étapes, des livrables inspectables",
        "L'extraction de métadonnées est distincte de l'export des données Hyper.",
        section, sources=("architecture",),
    )
    stages = (
        ("Sources Tableau", ".twb / .twbx", ORANGE),
        ("Extraction", "XML et métadonnées", INK),
        ("23 JSON", "Contrat intermédiaire", TEAL),
        ("Génération", "TMDL + PBIR + M", INK),
    )
    for i, (title, body, fill) in enumerate(stages):
        x = 0.65 + i * 3.1
        _node(slide, x, 2.83, 2.7, title, body, fill=fill)
        if i < 3:
            _arrow(slide, x + 2.79, 3.40)
    _panel(slide, 0.65, 4.67, 5.81, 1.0, "PBIP / PBIR v4.0 / TMDL", "",
           fill=WHITE)
    _panel(slide, 6.86, 4.67, 5.81, 1.0, "Sources + preuves de qualité", "",
           fill=WHITE)
    _banner(slide, "Générer n'est ni actualiser les données, ni certifier la parité, ni publier.")


def _source_policy(prs, section):
    slide = _slide(
        prs, "Base d’origine ou Hyper : choisir la source",
        "Même source Tableau : base explicite complète prioritaire ; sinon Hyper.",
        section, sources=("architecture",),
        note="Si une même source déclare Hyper et une base d’origine explicite, complète "
        "et non ambiguë, retenir cette base avant toute lecture/export Hyper. "
        "Ne pas deviner une base depuis Excel/CSV ou une référence publiée non résolue. "
        "Un lien de base incomplet ou ambigu est une erreur explicite. "
        "Hyper conservé <= 10 * 1024 * 1024 octets : CSV + Import. Strictement au-dessus : "
        "Parquet natif + Import, SDK Hyper optionnel requis. Base native : DirectQuery "
        "si le connecteur le permet, sinon Import. L'utilisateur peut forcer Import "
        "sans changer la source retenue. Aucune connexion à la base pendant la génération ; "
        "pas de garantie de parité avec un extrait filtré, agrégé ou daté. "
        "L'Import natif ne déclenche aucune extraction/copie de lignes par l'outil. "
        "Exception explicite : un classeur Excel réellement embarqué dans le TWBX est "
        "copié sans modification dans Data/ et lié par chemin absolu ; ce n'est pas une "
        "extraction de lignes, elle ne couvre ni les jobs téléversés ni les modèles "
        "partagés, et le chemin absolu local figure dans le projet. "
        "Fabric conserve un contrat de staging distinct.",
    )
    _tag(slide, 0.65, 2.53, 5.8, "1 — BASE D’ORIGINE SI ÉLIGIBLE", TEAL_LIGHT, TEAL)
    _tag(slide, 6.85, 2.53, 5.8, "2 — SINON : HYPER CONSERVÉ", ORANGE_LIGHT, ORANGE)
    _node(slide, 0.65, 3.19, 5.8, "DirectQuery si compatible",
          "Import si forcé ou si incompatible", fill=TEAL)
    _node(slide, 6.85, 3.19, 2.7, "≤ 10 MiB", "CSV + Import", fill="A3430C")
    _node(slide, 9.95, 3.19, 2.7, "> 10 MiB", "Parquet + Import", fill=INK)
    _text(slide, 0.72, 4.80, 5.6, 0.65,
          "Aucune lecture ni export du Hyper.\nConnexion, schéma et M ; accès à configurer.", 16)
    _text(slide, 6.92, 4.80, 5.6, 0.65,
          "Export complet des données Hyper.\nPour Parquet : SDK Hyper requis.", 16)
    _text(slide, 0.72, 5.57, 11.9, 0.28,
          "Exception — Excel embarqué : copié inchangé dans Data/, lié par chemin absolu.", 13, MUTED)
    _banner(slide, "Source retenue : parité avec l’extrait Tableau et actualisation restent à vérifier.")


def _evidence(prs, section):
    slide = _slide(
        prs, "Une preuve statique, pas une promesse de parité",
        f"Observations locales du {EVIDENCE_DATE} sur un corpus public / synthétique.",
        section, sources=("features", "authoring", "tests"),
        note="CLI Microsoft powerbi-report-author 0.5.0, validation avec schémas activés : "
        "4/10 avant, 22 erreurs ; 10/10 après, zéro erreur et zéro avertissement CLI. "
        "Qualité unifiée : 4 PASS, 6 WARN, 0 FAIL. Suite locale historique : 11 850 "
        "passed, 66 skipped, 1 xfailed, 1 warning. Ne pas confondre ces résultats "
        "avec la CI courante, le moteur DAX, Desktop, refresh ou des gains de temps.",
    )
    for x, number, title, body in (
        (0.65, "4 → 10 / 10", "Rapports validés", "Validation Microsoft PBIR"),
        (4.79, "22 → 0", "Erreurs CLI", "Schémas activés ; 0 warning final"),
        (8.93, "11 850", "Tests locaux réussis", "66 skipped ; 1 xfailed ; 1 warning"),
    ):
        _rect(slide, x, 2.78, 3.75, 2.10, WHITE, rounded=True)
        _text(slide, x + 0.23, 3.01, 3.29, 0.66, number, 35, TEAL, bold=True)
        _text(slide, x + 0.23, 3.85, 3.29, 0.33, title, 18, bold=True)
        _text(slide, x + 0.23, 4.32, 3.29, 0.36, body, 12, MUTED)
    _text(slide, 0.72, 5.27, 11.8, 0.36,
          "Qualité unifiée : 4 PASS / 6 WARN / 0 FAIL — les réserves restent visibles.",
          17, bold=True)
    _banner(slide, "Moteur modèle, ouverture Desktop, refresh, parité et déploiement : preuves séparées.")


def _delivery(prs, section):
    slide = _slide(
        prs, "Un workflow avec des décisions humaines",
        "L'outil prépare ; le responsable métier et la plateforme acceptent le résultat.",
        section, sources=("quality", "cli"),
    )
    labels = (
        ("Inventorier", "Sources, dépendances,\ncomplexité"),
        ("Générer", "Projet éditable,\nregistre des écarts"),
        ("Valider", "Refresh, valeurs,\nvisuels et sécurité"),
        ("Industrialiser", "Vagues, publication,\nsupport et reprise"),
    )
    for i, (title, body) in enumerate(labels):
        x = 0.65 + i * 3.1
        _rect_shape(slide, MSO_SHAPE.OVAL, x + 0.92, 2.70, 0.78, 0.78,
                    YELLOW if i == 2 else TEAL)
        _text(slide, x + 0.92, 2.86, 0.78, 0.43, str(i + 1), 24,
              INK if i == 2 else WHITE, bold=True, align=PP_ALIGN.CENTER)
        if i < 3:
            _arrow(slide, x + 2.57, 2.97, color=TEAL)
        _text(slide, x, 3.94, 2.7, 0.42, title, 22, bold=True)
        _text(slide, x, 4.56, 2.7, 0.94, body, 17, MUTED)
    _banner(slide, "Une génération réussie ne déclenche pas automatiquement une mise en production.")


def _sources(prs, section, *, business=False):
    slide = _slide(
        prs, "Sources, hypothèses et périmètre",
        "Les notes de chaque diapositive conservent les références et précautions utiles.",
        section, sources=("architecture", "features", "quality", "limits", "authoring"),
    )
    rows = (
        ("01", "Architecture & modes de source", "architecture"),
        ("02", "Fonctionnalités & preuves du 23.10.2026", "features"),
        ("03", "Checklist de recette & limites connues", "quality"),
        ("04", "Validation Microsoft optionnelle", "authoring"),
    )
    for i, (number, title, key) in enumerate(rows):
        y = 2.58 + i * 0.68
        _tag(slide, 0.65, y, 0.55, number)
        link = _text(slide, 1.48, y + 0.01, 10.9, 0.34, title, 18, BLUE)
        link.text_frame.paragraphs[0].runs[0].hyperlink.address = (
            f"{REPOSITORY}/blob/main/{REFERENCES[key][1]}"
        )
    _text(slide, 0.65, 5.62, 12.03, 0.72,
          "100 rapports ; 16 h manuelles ; 8 / 5 / 3 h assistées ; "
          "80 h de mise en place ; 100 €/h.\nHypothèses illustratives, non benchmarkées."
          if business else
          "La date d'édition n'est pas une version livrée.\n"
          "Les mappings ne garantissent ni la fidélité rendue ni la disponibilité des connecteurs.",
          15, MUTED)


def _overview(prs):
    _cover(prs)
    slide = _slide(
        prs, "Déplacer l'effort vers ce qui compte",
        "Automatiser la reconstruction répétitive, sans masquer le travail d'acceptation.",
        "Vision & valeur", sources=("features", "limits"),
    )
    for x, number, title, body in (
        (0.65, "01", "Réduire la répétition",
         "Extraire les métadonnées et générer un socle commun, plutôt que tout recréer à la main."),
        (4.79, "02", "Rendre les écarts visibles",
         "Afficher avertissements, approximations et vérifications non exécutées."),
        (8.93, "03", "Réutiliser le patrimoine",
         "Consolider les modèles et préparer des vagues cohérentes, avec arbitrage des conflits."),
    ):
        _panel(slide, x, 2.70, 3.75, 2.97, title, body, number=number)
    _banner(slide, "La valeur recherchée : de la capacité libérée et un processus plus reproductible.")

    _pipeline(prs, "Vision & valeur")
    slide = _slide(
        prs, "Six leviers fonctionnels",
        "Des capacités de génération et d'analyse ; pas une équivalence universelle.",
        "Fonctionnalités", sources=("features", "mapping", "limits"),
    )
    cards = (
        ("Évaluer", "Complexité, risques, dépendances et préparation des vagues."),
        ("Convertir", "Modèle TMDL, expressions DAX, requêtes M et rapport PBIR."),
        ("Consolider", "Modèle partagé et rapports fins ; conflits à arbitrer."),
        ("Préparer", "Prep vers M, sources et lignage ; périmètre borné."),
        ("Contrôler", "Rapports de qualité, preuves, écarts et reprises."),
        ("Accompagner", "Navigateur local, CLI et assistance IA optionnelle."),
    )
    for i, (title, body) in enumerate(cards):
        _panel(slide, 0.65 + (i % 3) * 4.14, 2.54 + (i // 3) * 1.87,
               3.75, 1.60, title, body)
    _source_policy(prs, "Stratégie des sources")
    _delivery(prs, "Modèle de delivery")
    _evidence(prs, "Preuves observées")

    slide = _slide(
        prs, "Comparer le même périmètre de travail",
        "Simulation commerciale ajustable — aucun benchmark de productivité mesuré.",
        "Hypothèses du business case",
        note="Les efforts assistés incluent revue humaine, remédiation et validation. "
        "80 heures de setup sont ajoutées une seule fois. Jours-personnes de 8 heures, "
        "pas un délai calendaire. Le taux chargé de 100 euros/heure valorise une capacité "
        "de travail et n'établit pas une économie de trésorerie.",
    )
    _panel(slide, 0.65, 2.66, 5.81, 2.81, "Référence manuelle",
           "100 rapports × 16 h\n= 1 600 heures\n= 200 jours-personnes", accent=ORANGE)
    _panel(slide, 6.86, 2.66, 5.81, 2.81, "Migration assistée",
           "80 h de mise en place\n+ 100 × effort assisté par rapport\nRevue et validation humaines incluses.")
    _banner(slide, "Taux chargé illustratif : 100 €/h  |  1 jour-personne = 8 h, pas un délai de livraison.")

    slide = _slide(
        prs, "Trois scénarios, pas une promesse",
        "Charge totale en heures ; les 80 h de mise en place sont incluses dans chaque scénario.",
        "Simulation illustrative",
        note="Référence : 1 600 h. Prudent : 80 + 100*8 = 880 h. "
        "Central : 80 + 100*5 = 580 h. Ambitieux : 80 + 100*3 = 380 h. "
        "Aucune de ces valeurs n'est un résultat de mesure terrain.",
    )
    baseline = SCENARIOS[0][1].manual_hours
    rows = (("Manuel", baseline, ORANGE),) + tuple(
        (name, case.assisted_hours, TEAL if name != "Central" else INK)
        for name, case in SCENARIOS
    )
    for i, (label, hours, fill) in enumerate(rows):
        y = 2.77 + i * 0.77
        _text(slide, 0.65, y + 0.03, 2.1, 0.40, label, 19, bold=True)
        _rect(slide, 3.02, y, 7.2 * hours / baseline, 0.44, fill, rounded=True)
        _text(slide, 10.60, y + 0.02, 2.0, 0.39,
              f"{_fr(hours)} h", 19, fill, bold=True)
    _banner(slide, "Effort assisté supposé : prudent 8 h / rapport ; central 5 h ; ambitieux 3 h.")

    slide = _slide(
        prs, "Valoriser la capacité potentiellement libérée",
        "Équivalent de charge évitée ; ni économie garantie, ni ROI net mesuré.",
        "Simulation illustrative",
        note="\n".join(
            f"{name}: manual={case.manual_hours}h, assisted={case.assisted_hours}h, "
            f"saved={case.saved_hours}h ({case.saved_percent}%), "
            f"person_days={case.saved_person_days}, labor_value={case.labor_value} EUR."
            for name, case in SCENARIOS
        ),
    )
    for i, (name, case) in enumerate(SCENARIOS):
        x = 0.65 + i * 4.14
        _rect(slide, x, 2.65, 3.75, 3.04, INK if i == 1 else WHITE, rounded=True)
        color = WHITE if i == 1 else INK
        accent = YELLOW if i == 1 else TEAL
        _text(slide, x + 0.23, 2.88, 3.29, 0.34, name, 20, color, bold=True)
        _text(slide, x + 0.23, 3.45, 3.29, 0.63,
              f"{_fr(case.saved_hours)} h", 37, accent, bold=True)
        _text(slide, x + 0.23, 4.23, 3.29, 0.40,
              f"{_fr(case.saved_percent, 2).removesuffix(',00')} % de charge évitée", 18, color)
        _text(slide, x + 0.23, 4.92, 3.29, 0.39,
              f"{_fr(case.labor_value)} € à 100 €/h", 18, color, bold=True)
    _banner(slide, "Cas central : 1 020 h = 127,5 jours-personnes de capacité, pas 127,5 jours de délai.")

    slide = _slide(
        prs, "Du gain brut au coût total du projet",
        "Le code ouvert n'annule ni les licences, ni l'infrastructure, ni le travail humain.",
        "Lecture économique",
        note="Le setup de 80 h et l'effort assisté incluent déjà la main-d'œuvre du scénario : "
        "ne pas les déduire une seconde fois. Définir C comme les coûts additionnels hors de "
        "cette main-d'œuvre (licences/capacité incrémentales, stockage, gateway, services "
        "non comptés, etc.). Valeur nette différentielle centrale = 102 000 EUR - C. "
        "ROI comparatif = (coût manuel - coût assisté total) / coût assisté total ; "
        "la référence, le périmètre et l'horizon doivent rester identiques.",
    )
    _rect(slide, 0.65, 2.63, 5.81, 2.94, INK, rounded=True)
    _text(slide, 0.90, 2.87, 5.31, 0.38,
          "Cas central illustratif", 20, WHITE, bold=True)
    _text(slide, 0.90, 3.30, 5.31, 0.38,
          "Capacité valorisée − coûts", 19, WHITE)
    _text(slide, 0.90, 3.83, 5.31, 0.66, "102 000 € − C", 38, YELLOW, bold=True)
    _text(slide, 0.90, 4.70, 5.31, 0.32,
          "C = coûts additionnels non déjà comptés.", 13, WHITE)
    _text(slide, 0.90, 5.13, 5.31, 0.30,
          "Ni économie de trésorerie garantie, ni ROI mesuré.", 13, WHITE)
    _text(slide, 6.86, 2.75, 5.81, 0.41, "À budgéter explicitement", 23, bold=True)
    for i, text in enumerate((
        "Licences Power BI / capacité Fabric",
        "Stockage, gateways et exploitation",
        "Intégration, formation et support",
        "Validation ou remédiation hors scénario",
    )):
        _tag(slide, 6.86, 3.43 + i * 0.55, 0.40, str(i + 1))
        _text(slide, 7.55, 3.44 + i * 0.55, 5.1, 0.37, text, 16)
    _banner(slide, "ROI comparatif = (coût manuel − coût assisté total) / coût assisté total ; même périmètre.")

    slide = _slide(
        prs, "Mesurer d'abord sur un pilote représentatif",
        "Calibrer les hypothèses avant toute extrapolation au portefeuille.",
        "Approche commerciale", sources=("quality", "limits"),
    )
    _panel(slide, 0.65, 2.65, 5.81, 2.96, "Mesurer la charge réelle",
           "Échantillon de 5 à 10 rapports variés.\nEffort : préparation, conversion,\n"
           "remédiation et recette.\nComparer le même périmètre.")
    _panel(slide, 6.86, 2.65, 5.81, 2.96, "Décider sur des critères acceptés",
           "Refresh et accès aux sources.\nTolérances métier sur les valeurs.\n"
           "Visuels, interactions et RLS.\nCoûts, dépendances et exceptions.")
    _banner(slide, "Sortie du pilote : effort observé, risques ouverts, TCO affiné, décision GO / NO GO.")

    slide = _slide(
        prs, "Livrer un patrimoine exploitable",
        "Un fichier généré ne suffit pas : la transmission fait partie du projet.",
        "Livrables & responsabilité", sources=("architecture", "quality", "deploy"),
    )
    for i, (title, body) in enumerate((
        ("Projets éditables", "PBIP, modèle TMDL,\nrapport PBIR et M"),
        ("Preuves & écarts", "Qualité, réserves,\ncontrôles non exécutés"),
        ("Sources & accès", "Bindings, exports Hyper,\ncredentials hors du dépôt"),
        ("Handoff", "Recette\nDéploiement autorisé\nRunbook et support"),
    )):
        _node(slide, 0.65 + i * 3.1, 2.95, 2.7, title, body, h=2.02,
              fill=TEAL if i == 3 else INK)
    _banner(slide, "La décision de mise en production reste séparée de la génération.")
    _sources(prs, "Références", business=True)
    slide = _slide(prs, "", "", "Passer à l'action", dark=True,
                   sources=("quality", "cli"))
    _text(slide, 0.65, 1.55, 11.6, 1.70, "Commencer petit.\nDécider sur des preuves.",
          44, WHITE, bold=True, font=FONT_TITLE, name="Closing title")
    _text(slide, 0.65, 3.80, 11.6, 0.57,
          "Pilote représentatif → mesure de l'effort → recette → migration par vagues.",
          21, LIGHT_TEXT)
    for i, label in enumerate(("CADRER", "PILOTER", "DÉCIDER", "INDUSTRIALISER")):
        _tag(slide, 0.65 + i * 3.1, 5.26, 2.7, label,
             YELLOW if i == 2 else NIGHT_CARD, INK if i == 2 else WHITE)


def _technical(prs):
    _cover(prs, technical=True)
    slide = _slide(
        prs, "Trois routes de sortie, trois contrats",
        "Le format dépend du besoin ; un Prep autonome n'est pas un rapport Power BI.",
        "Périmètre", sources=("architecture", "prep", "limits"),
    )
    _panel(slide, 0.65, 2.65, 3.75, 3.00, "Classeurs → PBIP",
           "TMDL + PBIR v4.0 + M.\nBatch de dossier, même pour un seul classeur.\n"
           "Modèle partagé : route séparée.", number=1)
    _panel(slide, 4.79, 2.65, 3.75, 3.00, "Prep → M & sources",
           "Power Query M, définitions de sources et assessment.\n"
           "Lignage inter-flows si au moins deux réussissent.", number=2)
    _panel(slide, 8.93, 2.65, 3.75, 3.00, "Fabric → artefacts",
           "Lakehouse, Dataflow Gen2, Notebook, DirectLake et Pipeline.\n"
           "Scaffold local, configuration puis déploiement distinct.", number=3)
    _banner(slide, "La génération directe d'un projet depuis un fichier positionnel est retirée.")
    _pipeline(prs, "Architecture logique")

    slide = _slide(
        prs, "Exécution locale, intégrations explicites",
        "Le navigateur loopback pilote le moteur ; les services distants ne sont pas implicites.",
        "Architecture d'exécution", sources=("architecture", "tools", "cli"),
    )
    for x, title, body in (
        (0.65, "Navigateur local", "127.0.0.1:8765"),
        (4.79, "CLI / automatisation", "migrate.py batch"),
        (8.93, "Surface agent / MCP", "Outils bornés et contrôlés"),
    ):
        _node(slide, x, 2.60, 3.75, title, body, fill=TEAL, h=1.20)
    _line(slide, 2.52, 3.80, 2.52, 4.13)
    _line(slide, 6.66, 3.80, 6.66, 4.13)
    _line(slide, 10.80, 3.80, 10.80, 4.13)
    _node(slide, 0.65, 4.13, 12.03, "Orchestration Python locale",
          "Extraction → génération → contrôles → fichiers et preuves sur disque", h=1.25)
    _text(slide, 0.70, 5.65, 11.9, 0.56,
          "Tableau Server, LLM, Desktop et publication cloud : intégrations séparées, "
          "avec prérequis et autorisations propres.", 16)
    _delivery(prs, "Workflow opérateur")
    _source_policy(prs, "Contrat des données")

    slide = _slide(
        prs, "PBIP : un projet lisible, pas une boîte noire",
        "Sources natives externes, sauf Excel embarqué copié dans Data/ ; l'Import charge à l'actualisation.",
        "Anatomie du résultat", sources=("architecture", "authoring"),
    )
    _rect(slide, 0.65, 2.60, 5.81, 3.22, INK, rounded=True)
    _text(slide, 0.93, 2.84, 5.25, 2.66,
          "Projet.pbip\n\nProjet.SemanticModel\\\n    definition\\  (TMDL)\n\n"
          "Projet.Report\\\n    definition\\  (PBIR)",
          18, WHITE, font="Cascadia Code")
    _panel(slide, 6.86, 2.60, 5.81, 1.43, "Modèle sémantique",
           "Tables, mesures, relations et partitions M.", fill=WHITE)
    _panel(slide, 6.86, 4.38, 5.81, 1.43, "Rapport",
           "Pages, visuels, filtres, signets et thème.", fill=WHITE)
    _banner(slide, "Configurer les accès aux sources, actualiser, sauvegarder puis rouvrir pour la recette.")

    slide = _slide(
        prs, "Le modèle sémantique porte la logique métier",
        "Les règles de conversion produisent un point de départ ; le contexte reste à vérifier.",
        "TMDL", sources=("features", "limits"),
    )
    cards = (
        ("Structure", "Tables, colonnes, types, descriptions et catégories."),
        ("Calculs", "Mesures DAX, colonnes calculées et dossiers d'affichage."),
        ("Relations", "Cardinalité et inférences à contrôler dans le modèle."),
        ("Temps", "Calendrier M, clés et tri des attributs temporels."),
        ("Sécurité", "Rôles RLS ; affectations et tests par identité distincts."),
        ("Navigation", "Hiérarchies, paramètres, groupes, sets et bins."),
    )
    for i, (title, body) in enumerate(cards):
        _panel(slide, 0.65 + (i % 3) * 4.14, 2.54 + (i // 3) * 1.87,
               3.75, 1.60, title, body)

    slide = _slide(
        prs, "PBIR : conserver l'intention visuelle",
        "164 enregistrements de mapping : un inventaire, pas 164 rendus certifiés.",
        "Rapports & interactions", sources=("mapping", "authoring", "limits"),
    )
    x = 0.65
    for count, label, fill in ((129, "natifs", TEAL), (16, "approximations", ORANGE),
                               (19, "custom", INK)):
        width = 12.03 * count / 164
        _rect(slide, x, 2.62, width, 0.51, fill)
        x += width
    for i, (number, label, fill) in enumerate((
        ("129", "Mappings natifs", TEAL),
        ("16", "Approximations", ORANGE),
        ("19", "Mappings custom", INK),
    )):
        x = 0.65 + i * 4.14
        _text(slide, x, 3.54, 1.3, 0.62, number, 34, fill, bold=True)
        _text(slide, x + 1.50, 3.69, 2.25, 0.40, label, 17, bold=True)
    _panel(slide, 0.65, 4.54, 5.81, 1.18, "Disposition & style",
           "Pages, thème, formatage et objets de dashboard.")
    _panel(slide, 6.86, 4.54, 5.81, 1.18, "Interaction",
           "Slicers, filtres, signets, navigation et tooltips.")
    _banner(slide, "Custom visuals, interactions et rendu exact : disponibilité et recette à confirmer.")

    slide = _slide(
        prs, "De la formule Tableau au bon contexte",
        "Classifier avant de convertir : une mesure et une colonne n'ont pas la même sémantique.",
        "DAX & Power Query M", sources=("features", "mapping", "limits"),
    )
    _node(slide, 0.65, 3.22, 2.70, "Formule Tableau", "Références et agrégations", fill=ORANGE)
    _arrow(slide, 3.46, 3.80)
    _node(slide, 3.85, 3.22, 2.70, "Classification", "Mesure ou colonne", fill=TEAL)
    _line(slide, 6.55, 3.91, 7.0, 3.91, TEAL)
    _line(slide, 7.0, 3.05, 7.0, 4.94, TEAL)
    _line(slide, 7.0, 3.05, 7.5, 3.05, TEAL)
    _line(slide, 7.0, 4.94, 7.5, 4.94, TEAL)
    _node(slide, 7.50, 2.52, 5.15, "Mesure → DAX", "Agrégation et contexte de filtre", h=1.25)
    _node(slide, 7.50, 4.42, 5.15, "Colonne → M si convertible",
          "Sinon DAX ; cas inter-tables à contrôler", h=1.25)
    _banner(slide, "133+ conversions référencées ; LOD et calculs de table peuvent nécessiter une reprise.")

    slide = _slide(
        prs, "Consolider sans effacer les conflits",
        "Une route de modèle partagé distincte du batch de projets indépendants.",
        "Multi-classeurs", sources=("architecture", "features", "limits"),
    )
    _node(slide, 0.65, 2.75, 2.70, "Classeurs A / B / C", "Schémas et calculs", fill=ORANGE)
    _arrow(slide, 3.46, 3.34)
    _node(slide, 3.85, 2.75, 4.15, "Matching + résolution",
          "Empreintes, colonnes, noms et conflits", fill=TEAL)
    _arrow(slide, 8.13, 3.34)
    _node(slide, 8.55, 2.75, 4.10, "Modèle commun", "Rapports fins par classeur")
    _panel(slide, 0.65, 4.67, 5.81, 1.08, "Décisions explicites", "Mesures, tables et paramètres.")
    _panel(slide, 6.86, 4.67, 5.81, 1.08, "Contrôles renforcés", "RLS, relations et sémantique métier.")
    _banner(slide, "Un score de rapprochement ne suffit pas à prouver l'identité métier de deux tables.")

    slide = _slide(
        prs, "Prep autonome : M, sources et lignage",
        "Pas de PBIP généré pour les fichiers .tfl / .tflx autonomes.",
        "Tableau Prep", sources=("prep", "architecture", "limits"),
    )
    for i, (title, body) in enumerate((
        ("Analyser le DAG", "Entrées, transformations, sorties"),
        ("Exporter", "PowerQuery/*.pq ; Sources/*.json"),
        ("Évaluer & relier", "assessment.json ; lignage inter-flows"),
    )):
        x = 0.65 + i * 4.14
        _node(slide, x, 2.77, 3.75, title, body, fill=TEAL if i == 1 else INK, h=1.61)
        if i < 2:
            _arrow(slide, x + 3.86, 3.43, w=0.17)
    _text(slide, 0.72, 4.93, 11.8, 0.71,
          "Lignage croisé automatique si au moins 2 flows réussissent.\n"
          "Parser et dialectes bornés : opérations non prises en charge à revoir.", 18)
    _banner(slide, "Joins, agrégations et expressions converties doivent être validés sur les données.")

    slide = _slide(
        prs, "Fabric : un scaffold à configurer et déployer",
        "Les artefacts sont générés localement ; la présence d'un fichier n'est pas une exécution cloud.",
        "Sortie Fabric", sources=("architecture", "deploy", "limits"),
    )
    for i, (title, body, fill) in enumerate((
        ("Dataflow Gen2", "Ingestion M", TEAL),
        ("Notebook", "ETL / transformations", INK),
        ("DirectLake", "Modèle sémantique", INK),
    )):
        x = 0.65 + i * 4.14
        _node(slide, x, 2.73, 3.75, title, body, fill=fill)
        if i < 2:
            _arrow(slide, x + 3.86, 3.32, w=0.17)
    _rect(slide, 0.65, 4.63, 12.03, 0.56, TEAL_LIGHT, rounded=True)
    _text(slide, 0.90, 4.76, 11.53, 0.29,
          "Lakehouse / OneLake : tables Delta et staging propres à la route Fabric", 17, TEAL, bold=True)
    _text(slide, 0.72, 5.48, 11.8, 0.38,
          "Pipeline : Dataflow → Notebook → refresh du modèle ; identifiants à renseigner.", 17)
    _banner(slide, "Connexions, capacité, identités, IDs d'activités et exécution restent à configurer.")

    slide = _slide(
        prs, "Empiler les preuves, sans les confondre",
        "Une vérification non exécutée reste not_run ; un PASS statique ne la remplace pas.",
        "Qualité & recette", sources=("quality", "authoring", "tools"),
    )
    stages = (
        ("Structure", "JSON, TMDL,\nréférences et schémas"),
        ("Modèle", "Moteur DAX,\nrelations et RLS"),
        ("Refresh", "Accès, données,\nsauvegarde / réouverture"),
        ("Parité", "Valeurs, captures,\ninteractions métier"),
    )
    for i, (title, body) in enumerate(stages):
        x = 0.65 + i * 3.1
        _node(slide, x, 2.74, 2.70, title, body, fill=TEAL if i == 0 else INK, h=1.92)
        if i < 3:
            _arrow(slide, x + 2.79, 3.57)
    _text(slide, 0.72, 5.16, 11.8, 0.50,
          "Puis validation opérationnelle et publication autorisée — chacune avec ses critères.", 18)
    _banner(slide, "Validation Microsoft PBIR optionnelle ≠ validation du moteur modèle ou du runtime Desktop.")
    _evidence(prs, "Preuves datées")

    slide = _slide(
        prs, "Réparer avec une trace et une nouvelle preuve",
        "Les réparations déterministes et les experts n'ont pas le même rôle.",
        "Auto-réparation & ownership", sources=("agents", "tools", "quality"),
    )
    for i, (title, body) in enumerate((
        ("Détecter", "Erreurs et écarts"),
        ("Corriger", "Règles déterministes"),
        ("Revalider", "Contrôles ciblés"),
        ("Tracer", "Preuves et réserves"),
    )):
        x = 0.65 + i * 3.1
        _node(slide, x, 2.68, 2.70, title, body, h=1.36, fill=TEAL if i == 2 else INK)
        if i < 3:
            _arrow(slide, x + 2.79, 3.25)
    _panel(slide, 0.65, 4.52, 5.81, 1.24, "17 spécialistes + 2 rôles de gouvernance",
           "Propriété par fichier ; handoff explicite.")
    _panel(slide, 6.86, 4.52, 5.81, 1.24, "Pas de dispatch automatique des experts",
           "Batch / shared-model n'invoquent pas le preceptor.")
    _banner(slide, "Une suggestion ou un ACK de handoff ne constitue pas une correction durable appliquée.")

    slide = _slide(
        prs, "Skills Microsoft, MCP et cœur déterministe",
        "Les références officielles améliorent l'authoring ; elles ne remplacent pas la recette.",
        "Assistance optionnelle", sources=("tools", "authoring", "agents"),
        note="L'intégration testée repose sur powerbi-report-author 0.5.0 et ses métadonnées. "
        "Le MCP de modélisation reste soumis à une EULA distincte non acceptée ici. "
        "Ne pas présenter une vérification de moteur comme exécutée. Les outils "
        "d'assistance et de handoff ne déclenchent pas automatiquement tous les agents.",
    )
    _panel(slide, 0.65, 2.65, 3.75, 2.95, "Références officielles",
           "skills-for-fabric.\nMétadonnées de rôles et formatage.\n"
           "CLI Microsoft de validation PBIR.", number=1)
    _panel(slide, 4.79, 2.65, 3.75, 2.95, "Outils & suggestions",
           "MCP borné et assistance LLM optionnelle.\n"
           "Contexte, masquage des données sensibles et contrôle humain.", number=2)
    _panel(slide, 8.93, 2.65, 3.75, 2.95, "Production vérifiable",
           "Générateurs déterministes.\nTests de non-régression.\n"
           "Écarts et preuves persistés.", number=3)
    _banner(slide, "EULA du MCP modèle non acceptée : aucun résultat d'exécution moteur n'est revendiqué.")

    slide = _slide(
        prs, "Séparer génération, accès et publication",
        "L'exploitation sécurisée dépend aussi des identités et de l'environnement cible.",
        "Gouvernance", sources=("features", "quality", "deploy"),
    )
    _panel(slide, 0.65, 2.65, 5.81, 3.00, "Protections dans le pipeline",
           "Validation des chemins et archives.\nParsing XML protégé.\n"
           "Détection / masquage des identifiants sensibles.\nJournaux et preuves à conserver.")
    _panel(slide, 6.86, 2.65, 5.81, 3.00, "Responsabilités d'exploitation",
           "Secrets hors dépôt et accès minimaux.\nIdentités, gateways et RLS testés.\n"
           "Politique de rétention des exports Hyper.\nApprobation explicite de publication.")
    _banner(slide, "Des contrôles intégrés ne remplacent pas une revue de sécurité propre au déploiement.")

    slide = _slide(
        prs, "Publier seulement après la recette",
        "La génération locale ne crée pas implicitement une mise en production.",
        "Déploiement", sources=("deploy", "quality"),
    )
    for i, (title, body) in enumerate((
        ("Configurer", "Workspace, identité,\nconnexions, gateway"),
        ("Prévisualiser", "Artefacts, dépendances,\npermissions et coûts"),
        ("Autoriser", "Déploiement séparé,\nacceptation du périmètre"),
        ("Vérifier", "Refresh, supervision,\nreprise et rollback"),
    )):
        x = 0.65 + i * 3.1
        _node(slide, x, 2.83, 2.70, title, body, h=2.0,
              fill=TEAL if i == 2 else INK)
        if i < 3:
            _arrow(slide, x + 2.79, 3.70)
    _banner(slide, "Licences Power BI / capacité Fabric et coûts d'exploitation restent à budgéter.")

    slide = _slide(
        prs, "Les commandes qui correspondent aux routes",
        "Exemples depuis la racine du dépôt ; dossiers source et sortie distincts.",
        "Démarrage", sources=("cli", "architecture"),
    )
    commands = (
        ("Navigateur local", r".\run_web_ui.ps1"),
        ("Batch, même pour 1 classeur",
         r"python migrate.py batch .\workbooks --output-dir .\output"),
        ("Forcer Import",
         r"python migrate.py batch .\workbooks --mode import --output-dir .\output"),
        ("Assessment sans génération", r"python migrate.py assess .\workbook.twbx"),
    )
    for i, (label, command) in enumerate(commands):
        y = 2.50 + i * 0.91
        _text(slide, 0.65, y, 12.03, 0.26, label, 13, TEAL, bold=True)
        _rect(slide, 0.65, y + 0.34, 12.03, 0.42, INK, rounded=True)
        _text(slide, 0.87, y + 0.42, 11.59, 0.25,
              command, 13, WHITE, font="Cascadia Code")

    slide = _slide(
        prs, "Un contrat de recette avant l'industrialisation",
        "Les approximations et prérequis deviennent des éléments de décision, pas des surprises.",
        "Critères d'acceptation", sources=("quality", "limits", "mapping"),
    )
    _panel(slide, 0.65, 2.65, 5.81, 2.98, "Vérifier sur un échantillon",
           "Sources représentatives et accès réels.\nRefresh, sauvegarde / réouverture.\n"
           "Valeurs aux tolérances convenues.\nCaptures comparées et interactions.")
    _panel(slide, 6.86, 2.65, 5.81, 2.98, "Documenter ce qui reste ouvert",
           "DAX / LOD, contexte et cardinalités.\nCustom visuals et connecteurs requis.\n"
           "Dialectes Prep non pris en charge.\nRLS, parité et contrôles not_run.")
    _banner(slide, "97 clés de dispatch, alias inclus → 55 fonctions de connecteurs distinctes ; pas 97 garanties.")
    _sources(prs, "Références techniques")
    slide = _slide(prs, "", "", "Synthèse technique", dark=True,
                   sources=("architecture", "quality", "deploy"))
    _text(slide, 0.65, 1.55, 11.6, 1.70, "Un socle généré.\nUne acceptation à démontrer.",
          44, WHITE, bold=True, font=FONT_TITLE, name="Closing title")
    _text(slide, 0.65, 3.80, 11.6, 0.76,
          "Métadonnées → projets éditables → preuves → décision humaine.",
          23, LIGHT_TEXT)
    for i, label in enumerate(("EXTRAIRE", "GÉNÉRER", "VALIDER", "AUTORISER")):
        _tag(slide, 0.65 + i * 3.1, 5.26, 2.7, label,
             YELLOW if i == 2 else NIGHT_CARD, INK if i == 2 else WHITE)


def _generate(output_path, *, technical):
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    prs.core_properties.title = (
        "Tableau to Power BI — Architecture et workflows"
        if technical else "Tableau to Power BI — Vision et business case"
    )
    prs.core_properties.subject = (
        f"Édition {EDITION_DATE} ; preuves datées ; hypothèses commerciales illustratives"
    )
    prs.core_properties.author = "TableauToPowerBI project"
    # The python-pptx template carries a third-party author and 2013 dates.
    prs.core_properties.last_modified_by = prs.core_properties.author
    prs.core_properties.created = EDITION_TIMESTAMP
    prs.core_properties.modified = EDITION_TIMESTAMP
    prs.core_properties.comments = "Generated by docs/generate_pptx.py"
    prs.core_properties.keywords = "Tableau, Power BI, PBIP, architecture, migration"
    (_technical if technical else _overview)(prs)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destination))
    print(f"Presentation saved: {destination}")
    print(f"  Slides: {len(prs.slides)}")
    return output_path


def generate_presentation(output_path):
    """Generate the overview, preserving the original entrypoint contract."""
    return _generate(output_path, technical=False)


def generate_technical_guide(output_path):
    """Generate the deeper architecture and operator guide."""
    return _generate(output_path, technical=True)


def main(default_deck="overview"):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--deck", choices=("overview", "guide", "both"), default=default_deck)
    output = parser.add_mutually_exclusive_group()
    output.add_argument("--output", "-o", type=Path, help="Output PPTX for a single deck")
    output.add_argument("--output-dir", type=Path, help="Directory for standard output names")
    args = parser.parse_args()
    if args.deck == "both" and args.output is not None:
        parser.error("--output is for a single deck; use --output-dir with --deck both")
    directory = args.output_dir if args.output_dir is not None else DOCS_DIR
    if args.deck in ("overview", "both"):
        generate_presentation(args.output if args.output is not None else directory / OVERVIEW_NAME)
    if args.deck in ("guide", "both"):
        generate_technical_guide(args.output if args.output is not None else directory / GUIDE_NAME)


if __name__ == "__main__":
    main()
